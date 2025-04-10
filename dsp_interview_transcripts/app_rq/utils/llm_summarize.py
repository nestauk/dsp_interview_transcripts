import ast
import os

from pathlib import Path
from typing import Dict
from typing import List
from typing import Tuple

import pandas as pd

from langchain.chains import LLMChain
from langchain.chat_models import AzureChatOpenAI
from langchain.prompts import PromptTemplate
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.worksheet.worksheet import Worksheet


llm = AzureChatOpenAI(
    openai_api_version=os.getenv("AZURE_OPENAI_API_VERSION"),
    azure_deployment=os.getenv("AZURE_OPENAI_DEPLOYMENT_NAME"),
    azure_endpoint=os.getenv("AZURE_OPENAI_ENDPOINT"),
    openai_api_key=os.getenv("AZURE_OPENAI_API_KEY"),
    temperature=0,
)

answer_prompt = PromptTemplate.from_template(
    "Based on the following documents:\n\n{context}\n\n"
    "Answer the question: '{question}'\n"
    "Give a concise summary answer based only on the information provided."
)

quote_prompt = PromptTemplate.from_template(
    "Here are some documents:\n\n{context}\n\n"
    "The answer to the question '{question}' was: {answer}\n"
    "Extract the 3 documents that best support this answer."
    "Return them as a valid Python list of 3 strings. Do not modify the selected documents in any way. No explanation.\n"
    'Format: ["Doc 1", "Doc 2", "Doc 3"]'
)


def summarize_and_quote(texts: List[str], question: str) -> Tuple[str, List[str]]:
    """
    Generate a concise summary answer and extract supporting quotes for a research question
    based on a list of input texts.

    Args:
        texts (List[str]): A list of textual excerpts (e.g., sentences or paragraphs)
            extracted from transcripts.
        question (str): The research question to answer based on the input texts.

    Returns:
        Tuple[str, List[str]]:
            - A concise summary answer to the research question.
            - A list of 3 direct quotes from the input texts that support the answer.
              If parsing fails, a fallback list with a single "[Parsing error]" entry is returned.
    """
    context = "\n\n".join(texts)
    answer_chain = LLMChain(llm=llm, prompt=answer_prompt)
    answer = answer_chain.run(context=context, question=question)

    quote_chain = LLMChain(llm=llm, prompt=quote_prompt)
    quotes = quote_chain.run(context=context, question=question, answer=answer)

    try:
        quotes_list = ast.literal_eval(quotes)
    except Exception:
        quotes_list = ["[Parsing error]"]
    return answer, quotes_list


def generate_single_summary(
    df: pd.DataFrame, question: str, rq_id: str, output_dir: str, test_mode: bool
) -> Tuple[pd.DataFrame, str, List[str]]:
    """
    Generate a summary and representative quotes for a single research question.

    Args:
        df (pd.DataFrame): Input DataFrame with columns including 'text' and 'identifier', where both are lists.
                        The fields 'text' and 'identifier' come from the way the LLM batch check is run in `llm_question_answering.py`.
        question (str): The research question to answer.
        rq_id (str): Unique ID for the research question, used for file naming. This comes from the way rq_dict is defined in the page `llm_analysis.py`.
        output_dir (str): Directory where intermediate outputs should be saved.
        test_mode (bool): If True, return a mock answer and a few sample quotes instead of calling the real summarizer.

    Returns:
        Tuple[pd.DataFrame, str, List[str]]: Exploded DataFrame, summary answer, and list of quotes.
    """
    df_long = df.explode(["text", "identifier"])
    df_long.to_csv(f"{output_dir}/{rq_id}_long.csv", index=False)

    extracted_texts = [txt for sublist in df["text"] for txt in sublist]

    if test_mode == True:
        answer = f"(TEST) This is a mock summary for: {question}"
        quotes = extracted_texts[:3]
    else:
        answer, quotes = summarize_and_quote(extracted_texts, question)

    return df_long, answer, quotes


def generate_summaries(
    rq_dict: Dict[str, str], output_paths: Dict[str, str], output_dir: str, test_mode: bool
) -> Tuple[Dict[str, Dict[str, object]], Dict[str, pd.DataFrame]]:
    """
    Iterate through RQS and generate a summary and key quotes for each one.

    Returns: a dictionary mapping each RQ to the summary answer and quotes;
    and a long-form dataframe of the batch_check output for each RQ.

    Args:
        rq_dict (Dict[str, str]): Dictionary mapping rq_id to question text.
        output_paths (Dict[str, str]): Dictionary mapping rq_id to path of corresponding JSONL file where the batch_check output (question answering) is saved.
        output_dir (str): Directory to save intermediate outputs in the Dash session.
        test_mode (bool): Whether to generate mock summaries or real ones.

    Returns:
        Tuple[
            Dict[str, Dict[str, object]],  # per_rq_outputs
            Dict[str, pd.DataFrame]        # long_dfs
        ]
    """
    per_rq_outputs = {}
    long_dfs = {}
    for rq_id, question in rq_dict.items():

        output = {}

        path = Path(output_paths[rq_id])

        df = pd.read_json(path, lines=True)

        long_dfs[question], output["answer"], output["quotes"] = generate_single_summary(
            df, question, rq_id, output_dir, test_mode
        )

        per_rq_outputs[question] = output

    return per_rq_outputs, long_dfs


def generate_full_summary_output(
    rq_dict: Dict[str, str], long_dfs: Dict[str, pd.DataFrame], per_rq_outputs: Dict[str, Dict[str, object]]
) -> pd.DataFrame:
    """
    Merge answers and quotes for each research question generated by summarize_and_quote() with the output produced at the question answering stage.
    This is used to make sure we have a UUID and conversation ID for each of the selected quotes.

    Args:
        rq_dict (Dict[str, str]): Dictionary mapping rq_id to question text.
        long_dfs (Dict[str, pd.DataFrame]): Exploded DataFrames for each question.
        per_rq_outputs (Dict[str, Dict[str, object]]): Output summaries and quotes for each question.

    Returns:
        pd.DataFrame: Combined summary DataFrame with merged quotes and answers.
    """
    full_summary_df = pd.DataFrame()

    for _, question in rq_dict.items():
        merged_output = pd.merge(
            long_dfs[question][["text", "identifier", "id"]],
            pd.DataFrame(per_rq_outputs[question]),
            left_on="text",
            right_on="quotes",
            how="right",
        )

        merged_output["question"] = question

        full_summary_df = pd.concat(
            [full_summary_df, merged_output[["question", "answer", "id", "identifier", "text", "quotes"]]],
            ignore_index=True,
        )
    return full_summary_df


def merge_column(ws: Worksheet, col_idx: int, start_row: int = 2) -> None:
    """
    Merge rows within a column that have the same value.

    The start row is 2 because row 1 is assumed to be the header.
    """
    row = start_row
    while row <= ws.max_row:
        current_value = ws.cell(row=row, column=col_idx).value
        end_row = row
        while end_row + 1 <= ws.max_row and ws.cell(row=end_row + 1, column=col_idx).value == current_value:
            end_row += 1
        if end_row > row:
            ws.merge_cells(start_row=row, start_column=col_idx, end_row=end_row, end_column=col_idx)
            ws.cell(row=row, column=col_idx).alignment = Alignment(vertical="center")
        row = end_row + 1


def create_output_excel(full_summary_df: pd.DataFrame, output_dir: str) -> None:
    """
    Save the full output (answers to RQs and illustrative quotes) to excel.
    This excel file is then available for download in the Dash session.

    Args:
        full_summary_df (pd.DataFrame): The full summary output DataFrame.
        output_dir (str): Directory to save the Excel file to.
    """
    df = full_summary_df.copy()

    # List of columns to merge
    columns_to_merge = ["question", "answer"]

    # Save DataFrame to Excel
    output_path = f"{output_dir}/full_summary.xlsx"
    df.to_excel(output_path, index=False)

    # Load with openpyxl
    wb = load_workbook(output_path)
    ws = wb.active

    # Apply merge to each column in `columns_to_merge`
    for col_name in columns_to_merge:
        col_idx = list(df.columns).index(col_name) + 1  # Convert to Excel's 1-indexing
        merge_column(ws, col_idx)

    # Save the result
    wb.save(output_path)
